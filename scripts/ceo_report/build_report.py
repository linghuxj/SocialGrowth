from pathlib import Path
from decimal import Decimal as D
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('/Users/linghuxj/Documents/myproject/project/SocialGrowth')
OUT=ROOT/'artifacts/reports/SocialGrowth年度技术规划与预算汇报.docx'
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.footer_distance=Cm(.8)
for name,size in [('Normal',10.5),('Title',23),('Subtitle',11),('Heading 1',15),('Heading 2',12)]:
 st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Heiti SC')
 st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.2
 if name.startswith('Heading'):st.paragraph_format.space_before=Pt(12)
doc.styles['Normal'].paragraph_format.widow_control=True
# Page numbering is useful for the six-page executive report.
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('第 ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
p.add_run(' 页').font.size=Pt(9)
doc.core_properties.title='SocialGrowth年度技术规划与预算汇报'
doc.core_properties.subject='年度技术投入 业务扩展及月度交付'
doc.core_properties.author='项目技术团队'
doc.core_properties.keywords='SocialGrowth 年度规划 技术预算 月度交付'

def para(text,bold=False,style=None):
 p=doc.add_paragraph(style=style);r=p.add_run(text);r.bold=bold;return p

def heading(text,level=1):
 p=doc.add_heading(text,level);p.paragraph_format.space_before=Pt(16 if level==2 else 12)
def page():doc.add_page_break()
def table(headers,rows,widths,center_cols=(),font=10):
 t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 pr=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for name in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr.append(borders)
 for idx,vals in enumerate([headers]+rows):
  row=t.rows[0] if idx==0 else t.add_row()
  rp=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');rp.append(ns)
  if idx==0:
   rep=OxmlElement('w:tblHeader');rp.append(rep)
  for j,txt in enumerate(vals):
   c=row.cells[j];c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1' if idx==0 else ('F5F7F9' if idx%2==0 else 'FFFFFF'));tcp.append(sh)
   mar=OxmlElement('w:tcMar')
   for key,val in [('top','85'),('bottom','85'),('left','100'),('right','100')]:
    e=OxmlElement('w:'+key);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');mar.append(e)
   tcp.append(mar)
   p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.16
   if j in center_cols:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   r=p.add_run(str(txt));r.font.size=Pt(font);r.bold=(idx==0);r.font.color.rgb=RGBColor(0,0,0)
 # spacer
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1;p.add_run().font.size=Pt(3)
 return t

# PAGE 1
para('SocialGrowth年度技术规划与预算汇报',style='Title')
para('汇报对象  CEO    汇报单位  项目技术团队',style='Subtitle')
para('规划周期  项目启动后十二个月    日期  2026年9月8日',style='Subtitle')
heading('一 汇报摘要')
para('拟建设面向企业的 AI 社交媒体运营与增长系统，以 AI 漫剧切片运营验证内容匹配、策略决策、自动分发和效果反馈能力；第六个月拓展第二条业务线，第十个月形成四条业务并行运行或试点，第十二个月完成同类客户接入验证及年度验收。')
para('年度技术预算为 162–196 万元，覆盖三人团队、AI 研发工具、基础设施、产品 AI 调用和软件服务五类投入。团队全年配置保持三人，模型调用、基础设施和软件服务随业务规模分阶段增加。',bold=True)
table(['核心事项','年度规划'],[
 ['人员配置','技术负责人、研发、产品测试，共三人；人力合计 8 万元/月。'],
 ['首期平台','Facebook、Instagram、YouTube 同步验证，首月每平台 10 个账号；第二月累计试点 90 个账号。'],
 ['账号增长','第 4–7 月累计万粉账号目标依次为 2–3、5–7、9–13、15–20 个。'],
 ['业务扩展','第六个月一个新业务场景完成真实试点；第十个月新增两条，总计四条业务。'],
 ['年度成果','四条业务持续运营，已验证场景自主处理，至少一个同类客户真实接入验证。']],[3,14])
heading('二 业务定位与服务范围')
para('系统服务 MCN 及有出海内容增长需求的企业，按地区、语言、人群和垂直领域管理账号，将适合的短视频内容匹配到对应账号。AI 负责分析、策略拟定及已验证场景的执行，人员负责业务确认、异常处理和效果核对。')
para('业务范围为内容生产与加工、账号运营、导流及效果分析。甲方负责承接平台；实物电商场景下，甲方负责商品、店铺、交易、库存、物流与售后。新业务线可包括出海实物电商品类。')

# PAGE 2
page();heading('三 人员配置与费用用途')
para('三人团队按每月 8 万元统一核算，全年 96 万元，包含公司完整用人成本，不按个人拆分预算。各岗位职责如下。')
table(['岗位','主要职责','交付责任'],[
 ['技术负责人','技术架构、AI 决策与规则设计、研发统筹、关键问题处理。','系统方案、关键技术评审、集成与技术验收。'],
 ['研发','业务系统开发、接口集成、任务执行与数据采集、技术验证和维护。','可运行版本、平台与设备对接、运行修复。'],
 ['产品测试','需求梳理、对接沟通、运营测试、统计指标核对及部分账号处理。','需求与验收清单、业务测试记录、运营反馈；暂不参与代码开发。']],[2.6,7.1,7.3])
heading('四 工具与运行资源用途')
table(['费用类别','具体作用','核算范围'],[
 ['AI 研发工具','Codex 与 Claude Code 用于辅助开发、代码理解、调试和测试，提升研发交付效率。','团队合计 1 万元/月。服务研发过程，与线上产品 AI 调用分别核算。'],
 ['基础设施','承载业务系统与任务调度，提供数据库、素材存储、流量、日志、监控和备份。','按共享资源统一配置，月预算从 1 万元提高至 1.5 万元、2 万元。'],
 ['产品 AI 调用','分析视频与运营数据、生成策略和调整方案，并按需完成文案、字幕、配音、图片及视频生成。','按实际产品任务消耗计费，月额度随阶段从 2–5 万元提高至 3–5 万元、4–7 万元。'],
 ['软件服务','邮件通知与对接、第三方统计分析、数据查询及相关接口服务，支撑运营协同与数据补充。','全年 6 万元，月均 5,000 元；前期较少，后期增加。与基础设施、AI 调用及研发工具分别核算。']],[2.7,7.4,6.9],font=10)
para('设备端已有自动化执行能力，技术工作以接口对接、任务下发和执行反馈为主。软件服务中的统计与查询能力按实际可用数据接入，避免同一订阅或接口费用在多类预算中重复计入。')

# PAGE 3
page();heading('五 年度预算及投入节奏')
para('以下金额均为人民币万元。预算下沿对应较低调用负载，上沿覆盖较高模型及内容生产消耗；实际支出按月统计。后期产品 AI 额度已包含按需视频生成，不再额外叠加生成预算。')
table(['预算类别','年度金额','计算口径'],[
 ['三人团队','96','8 × 12 个月'],['AI 研发工具','12','1 × 12 个月'],
 ['基础设施','16','1 × 7 个月 + 1.5 × 2 个月 + 2 × 3 个月'],
 ['产品 AI 调用','32–66','第 1–7 月 2–5/月；第 8–9 月 3–5/月；第 10–12 月 4–7/月'],
 ['软件服务','6','第 1–4 月 0.3/月；第 5–8 月 0.5/月；第 9–12 月 0.7/月'],
 ['年度合计','162–196','五类技术投入合并核算']],[3.3,2.4,11.3],center_cols=(1,))
heading('月度预算明细',2)
rows=[]
for m in range(1,13):
 infra=D('1') if m<=7 else D('1.5') if m<=9 else D('2')
 lo,hi=(D(2),D(5)) if m<=7 else (D(3),D(5)) if m<=9 else (D(4),D(7))
 soft=D('.3') if m<=4 else D('.5') if m<=8 else D('.7')
 total_lo=D(9)+infra+lo+soft;total_hi=D(9)+infra+hi+soft
 rows.append((m,infra,lo,hi,soft,total_lo,total_hi))
def f(x):return format(x,'f').rstrip('0').rstrip('.') if '.' in format(x,'f') else str(x)
tr=[[f'第{m}月','8','1',f(infra),f'{f(lo)}–{f(hi)}',f(soft),f'{f(tlo)}–{f(thi)}'] for m,infra,lo,hi,soft,tlo,thi in rows]
tr.append(['全年','96','12','16','32–66','6','162–196'])
table(['月份','人员','工具','设施','AI调用','软件服务','合计'],tr,[1.65,1.6,1.6,1.6,2.65,2.1,5.8],center_cols=tuple(range(7)),font=9.5)
para('软件服务按前低后高排期，全年平均每月 0.5 万元。上述预算覆盖列示五类费用；新增硬件采购及独立专项费用按实际需求另行核算。')

# PAGE 4
page();heading('六 月度交付清单')
para('第 1–6 月以漫剧业务建立可运行、可验证的增长流程，并开始新业务试点。所有目标按项目启动后的月份计，不对应固定自然月。')
table(['月份','主要交付与业务目标','验收材料'],[
 ['第1月','四周内三平台分别跑通策略确认、设备执行、短链导流和数据反馈；各 10 个初始账号。','三平台真实发布与执行记录、指标回收记录、短链统计；首期以已有切片为主。'],
 ['第2月','持续运营并开放已验证场景自动决策；每平台新增 20 个，累计试点 90 个。','连续运行、人工干预与异常处理、策略调整、效果及成本记录。'],
 ['第3月','形成可验证的策略优化流程，AI 根据效果提出改进并分批验证。','候选策略版本、对照记录、实验结果及保留或回退依据。'],
 ['第4月','形成 2–3 个标准万粉账号，持续验证内容匹配与运营策略。','账号名单、粉丝或订阅数据、运营状态与策略效果记录。'],
 ['第5月','累计形成 5–7 个万粉账号，扩大有效策略应用范围。','新增及退出达标记录、内容效果、人工投入与成本复盘。'],
 ['第6月','累计形成 9–13 个万粉账号；一个新业务场景完成首个真实运营试点。','原业务运行记录；新业务资料与规则、真实发布、导流反馈及首轮复盘。']],[1.95,7.85,7.2],font=10.5)
heading('关键业务指标',2)
para('重点跟踪观看量、可取得的观看完成情况、互动数、粉丝增长与短链点击。万粉账号指单账号粉丝或订阅数达到 10,000 以上；累计形成数量与当月仍达标、可运营数量分别展示，并记录掉粉、受限和恢复情况。')
para('账号数量是运营资产规模指标。增长质量同时结合内容表现、受众匹配和导流效果判断，不能只用粉丝数替代业务效果。')

# PAGE 5
page();heading('七 多业务扩展及年度交付')
para('第 7–12 月复用已有系统扩展业务，形成四条业务并行运营或试点，并验证同类客户接入能力。各业务分别配置内容规则、策略、导流目标和指标。')
table(['月份','主要交付与业务目标','验收材料'],[
 ['第7月','累计形成 15–20 个万粉账号；两业务持续运营，完成新业务策略优化。','两线运行记录、新业务实验报告及可复用的同类客户接入配置。'],
 ['第8月','完善两业务的资料输入、内容规则、策略及效果和费用统计的配置复用。','可配置流程、复用结果、逐业务运行及费用报告。'],
 ['第9月','确定新增两条业务场景，完成资料、内容规则、账号执行和数据接口适配。','两条新业务的适配验证记录及可进入真实试点的版本。'],
 ['第10月','原两条持续运营，新增两条分别完成真实试点；总计四条业务进入运行或试点。','四条业务分别提供真实执行、反馈、效果、成本和人工干预记录。'],
 ['第11月','四条业务持续运营，已验证场景由 AI 自主处理。','逐业务运行、策略优化、人工耗时、单位任务成本及异常恢复验证。'],
 ['第12月','至少一个新增同类客户完成真实接入验证，完成年度验收。','客户接入与真实闭环记录、配置模板、操作说明及四业务年度报告。']],[1.95,7.85,7.2],font=10.5)
heading('复制能力的验收方式',2)
para('第十二个月从一条已验证业务中选择同类客户，利用现有模板和配置完成资料接入、内容策略、真实发布和数据反馈。记录接入时间、人工投入及新增适配工作，判断系统复用是否减少重复开发。该客户接入属于既有业务复制。')
para('第十个月四条业务进入运行或试点，第十一、十二个月进一步验证持续运营和复用能力；各阶段以实际运行证据验收。')

# PAGE 6
page();heading('八 运营机制与管理检查清单')
heading('决策与执行流程',2)
para('客户目标与素材输入后，系统按账号定位匹配内容，结合平台规则、人工经验和运营数据生成策略。策略经约定审核后下发执行，系统回收发布结果、社媒数据和短链点击，形成下一轮优化。已验证场景逐步由 AI 自主处理，未知异常、规则冲突及重大调整由人员介入。')
heading('月度管理检查',2)
table(['检查项','汇报内容','管理用途'],[
 ['交付完成情况','当月系统版本、真实闭环记录及未完成事项。','核对月度目标与实际结果，明确后续工作。'],
 ['业务效果','按平台、业务和账号呈现观看、互动、粉丝及短链点击。','识别有效场景，确定扩大或调整策略。'],
 ['自动化与稳定性','自动处理范围、人工干预次数和耗时、任务异常及恢复结果。','判断团队承载能力和运行稳定性。'],
 ['费用消耗','人员、工具、设施、产品 AI、软件服务分别列预算与实际；模型消耗按业务归集。','定位费用增长来源，控制重复订阅与无效调用。'],
 ['业务复用','资料与规则复用情况、新客户接入耗时及必要适配工作。','评估新业务和同类客户扩展效率。']],[3.1,8.1,5.8],font=10.5)
heading('数据与验收口径',2)
para('短链统计反映导流点击。甲方平台内的注册、观看、下单等转化，仅在甲方授权并能够取得数据时纳入。不同平台的指标按各自可用定义记录，保留数据来源与采集时间。')
para('真实试点应具备客户资料、可执行账号和数据访问条件。阶段验收逐业务核对持续运行、有效样本、异常处理和人工干预情况；具体阈值纳入实施验收清单，业务目标以真实数据复盘。')
heading('提请审阅事项',2)
para('提请审阅十二个月交付计划、三人团队配置及五类技术投入安排。执行中按月提交交付和费用报告，依据业务结果、资源负载及模型消耗调整投入节奏；第六个月和第十个月分别结合新业务试点进行阶段评估。')

assert sum(r[4] for r in rows)==D(6)
assert sum(r[5] for r in rows)==D(162)
assert sum(r[6] for r in rows)==D(196)
# Remove template theme and border residue and explicitly set all script fonts.
for element in [doc.styles.element,doc.element,sec.footer._element]:
 for border in list(element.iter(qn('w:pBdr'))): border.getparent().remove(border)
 for fonts in element.iter(qn('w:rFonts')):
  for attr in list(fonts.attrib):
   if 'theme' in attr.lower(): del fonts.attrib[attr]
  for key in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+key),'Heiti SC')
for part in [doc.element,sec.footer._element]:
 for run in part.iter(qn('w:r')):
  rp=run.find(qn('w:rPr'))
  if rp is None: rp=OxmlElement('w:rPr');run.insert(0,rp)
  rf=rp.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts');rp.insert(0,rf)
  for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Heiti SC')
doc.styles['Subtitle'].font.italic=False
doc.save(OUT)
print(str(OUT))
print('Budget checked: 162–196 万元; software 6 万元; 12 monthly rows')
